from __future__ import annotations

import importlib.util
import os
from pathlib import Path

import openpyxl

from conftest import AGNI

EXAMPLE = Path(__file__).resolve().parents[1] / "examples" / "audit_workbook.py"


def test_audit_workbook_example_runs(tmp_path, isolated, monkeypatch):
    env, _ = isolated
    for k in ("HOME", "XDG_CONFIG_HOME"):
        monkeypatch.setenv(k, env[k])
    spec = importlib.util.spec_from_file_location("audit_workbook", EXAMPLE)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)

    out = tmp_path / "audit.xlsx"
    mod.main(["-o", str(out), "--agni", AGNI])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == [name for name, _ in mod.TABLES] + ["Findings"]
    # Every sheet carries rows on the tutorial board, so a query that silently matched nothing fails.
    for ws in wb.worksheets:
        assert ws.max_row > 1, f"{ws.title} is empty"


REVISION_AUDIT = EXAMPLE.with_name("revision_audit.py")
# The summary the example writes, committed so the gate regenerates and compares it (agni issue 823).
# AGNI_UPDATE_GOLDEN=1 rewrites it; read the diff before committing it, since it is the example's
# whole visible output.
REVISION_AUDIT_SUMMARY = EXAMPLE.with_name("revision_audit.summary.md")


def _run_revision_audit(tmp_path, isolated, monkeypatch):
    env, _ = isolated
    for k in ("HOME", "XDG_CONFIG_HOME"):
        monkeypatch.setenv(k, env[k])
    spec = importlib.util.spec_from_file_location("revision_audit", REVISION_AUDIT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    out = tmp_path / "audit.xlsx"
    mod.main(["-o", str(out), "--agni", AGNI])
    return openpyxl.load_workbook(out), out.with_suffix(".md").read_text()


def _rows(ws):
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def test_revision_audit_example_builds_every_tab(tmp_path, isolated, monkeypatch):
    wb, summary = _run_revision_audit(tmp_path, isolated, monkeypatch)
    audit = ["Pin to net", "Part numbers", "Test points per net", "One test point away", "Part numbers never probed", "PMIC rails"]
    built = ["Diff"] + audit + ["Review", "Review summary", "Findings", "Skipped rules", "Verdicts", "Verdicts by rule"]
    assert wb.sheetnames == built
    for ws in wb.worksheets:
        assert ws.max_row > 1, f"{ws.title} is empty"
    for name in built:
        assert f"## {name}\n" in summary, f"summary has no section for {name}"

    # A few rows a reader of the tutorial can check by eye.
    diff = _rows(wb["Diff"])
    assert ["net-renamed", "CLK_IN", "XTAL_IN"] in [r[:3] for r in diff]
    assert sum(1 for r in diff if r[0] == "net-equal") == 10
    # Uncovered nets answer 0 rather than vanishing (agni issue 819).
    assert ["GND", "0"] in [r[:2] for r in _rows(wb["Test points per net"])]
    # A pin designator orders as a name, so U3's pin 2 comes before its pin 10.
    u3 = [r[1] for r in _rows(wb["Pin to net"]) if r[0] == "U3"]
    assert u3.index("2") < u3.index("10")
    assert _rows(wb["Review"])[0][1] == "P1"
    # Rev B is read with its own board, so its copper rules ran (agni issue 848).
    assert any(r[2] == "copper-clearance" for r in _rows(wb["Findings"]))


def test_revision_audit_summary_is_the_committed_one(tmp_path, isolated, monkeypatch):
    _, summary = _run_revision_audit(tmp_path, isolated, monkeypatch)
    if os.environ.get("AGNI_UPDATE_GOLDEN"):
        REVISION_AUDIT_SUMMARY.write_text(summary)
    assert summary == REVISION_AUDIT_SUMMARY.read_text(), (
        "the revision-audit summary changed; rerun with AGNI_UPDATE_GOLDEN=1, read the diff, and commit it"
    )


# The example is the claim that an audit is data plus a short script, so both halves are held.
def test_revision_audit_script_stays_short_and_judgement_free():
    source = REVISION_AUDIT.read_text()
    assert len(source.splitlines()) < 200
    assert "=>" not in source, "a query belongs in revision_audit.yaml, not in the script"
