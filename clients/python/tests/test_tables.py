from __future__ import annotations

import openpyxl
import pytest

from agni import rows_as_dicts, table_sheets, tables_to_xlsx, to_rows
from agni.v1.checks import checks_pb2 as checks
from agni.v1.webapi import checks_pb2, query_pb2, tables_pb2


def _answer() -> query_pb2.RunQueryResponse:
    return query_pb2.RunQueryResponse(
        columns=["c", "n"],
        rows=[query_pb2.QueryRow(cells=["R1", "10"]), query_pb2.QueryRow(cells=["R2", "007"])],
    )


def test_rows_as_dicts_keeps_cells_as_sent():
    assert rows_as_dicts(_answer()) == [{"c": "R1", "n": "10"}, {"c": "R2", "n": "007"}]


def _table(name="t") -> tables_pb2.Table:
    return tables_pb2.Table(
        name=name,
        columns=[tables_pb2.TableColumn(name="c"), tables_pb2.TableColumn(name="n")],
        rows=[tables_pb2.TableRow(cells=["R1", "10"]), tables_pb2.TableRow(cells=["R2", "007"])],
    )


def test_to_rows_reads_an_engine_table_and_refuses_an_answer():
    assert to_rows(_table()) == (["c", "n"], [["R1", "10"], ["R2", "007"]])
    with pytest.raises(TypeError, match="tabulate"):
        to_rows(checks_pb2.CheckDesignResponse())


def test_table_sheets_keep_order_and_rename():
    resp = tables_pb2.TabulateResponse(tables=[_table("findings"), _table("verdicts")])
    assert [n for n, _ in table_sheets(resp, {"findings": "Findings"})] == ["Findings", "verdicts"]


def test_xlsx_has_header_freeze_and_filter(tmp_path):
    out = tmp_path / "audit.xlsx"
    tables_to_xlsx(str(out), [("Resistors", _table()), ("Plain", (["a"], [["1"], ["2"], ["3"]]))])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Resistors", "Plain"]
    ws = wb["Resistors"]
    assert [c.value for c in ws[1]] == ["c", "n"]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:B3"
    # "007" stays a string: a cell is what agni sent, not a guess at its type.
    assert ws["B3"].value == "007"
    assert wb["Plain"].auto_filter.ref == "A1:A4"


@pytest.mark.parametrize("name", ["", "a" * 32, "a/b", "x[1]"])
def test_xlsx_refuses_a_sheet_name_excel_rejects(tmp_path, name):
    with pytest.raises(ValueError):
        tables_to_xlsx(str(tmp_path / "x.xlsx"), [(name, _table())])


def test_xlsx_refuses_a_name_used_twice(tmp_path):
    with pytest.raises(ValueError, match="twice"):
        tables_to_xlsx(str(tmp_path / "x.xlsx"), [("Nets", _table()), ("nets", _table())])


def test_xlsx_keeps_a_formula_shaped_cell_as_text(tmp_path):
    out = tmp_path / "f.xlsx"
    tables_to_xlsx(str(out), [("Nets", (["net"], [["=HYPERLINK(\"x\")"], ["+5V"]]))])
    ws = openpyxl.load_workbook(out)["Nets"]
    assert ws["A2"].data_type == "s" and ws["A2"].value == '=HYPERLINK("x")'
    assert ws["A3"].value == "+5V"


# ---- natural order, and the diff, review and verdict sheets (agni issue 822) ----

from agni.v1.webapi import diff_pb2, review_pb2  # noqa: E402

from conftest import DESIGN  # noqa: E402


@pytest.fixture(params=["cli", "connect"])
def client(request):
    return request.getfixturevalue(request.param)


def _by_name(resp):
    return {t.name: to_rows(t) for t in resp.tables}


def test_tabulate_a_diff_keeps_near_rename_evidence(client, tmp_path):
    diff = client.diff_designs(a_uri=DESIGN + "/gateway-rev-b.edn", b_uri=DESIGN + "/gateway-rev-c.edn", near_renames={})
    out = client.tabulate(diff=diff, order_by=["change_class", "subject"])
    header, rows = _by_name(out)["diff"]
    # The header `agni diff --format csv` publishes: one table, change_class naming each row's kind.
    assert header[:3] == ["change_class", "subject", "old_name"]
    near = [dict(zip(header, r)) for r in rows if r[0] == "net-renamed-approx"]
    assert [(n["subject"], n["old_name"]) for n in near] == [("PMIC_ENABLE", "PMIC_EN")]
    assert near[0]["match_old_coverage"] and near[0]["added"] == "R6.1"
    tables_to_xlsx(str(tmp_path / "diff.xlsx"), table_sheets(out))


def test_tabulate_a_diff_carries_unchanged_nets_when_asked(client):
    diff = client.diff_designs(a_uri=DESIGN + "/gateway.edn", b_uri=DESIGN + "/gateway-rev-b.edn", include_equal=True)
    _, rows = _by_name(client.tabulate(diff=diff))["diff"]
    assert sum(1 for r in rows if r[0] == "net-equal") == 10
    assert [r[1] for r in rows if r[0] == "net-renamed"] == ["CLK_IN", "CLK_OUT"]


_CHECKLIST = checks.ReviewManifest(
    name="tables",
    areas=[
        checks.ManifestArea(
            name="Power",
            items=[
                checks.ManifestItem(id="P1", title="rails carry bulk capacitance", binding=checks.ItemBinding(rule="bulk-cap")),
                checks.ManifestItem(id="P2", title="i2c pull-ups", binding=checks.ItemBinding(rule="i2c-pull-up")),
                checks.ManifestItem(id="P3", title="reviewed by hand", note="the lead signs this off"),
            ],
        )
    ],
)


def test_tabulate_a_review_keeps_checklist_order(client, tmp_path):
    out = client.tabulate(review=client.create_review(design_uri=DESIGN, manifest=_CHECKLIST))
    tables = _by_name(out)
    header, rows = tables["review"]
    assert header == ["area", "id", "title", "outcome", "note", "findings", "unmet"]
    assert [r[1] for r in rows] == ["P1", "P2", "P3"]
    failed = [dict(zip(header, r)) for r in rows if r[3] == "fail"]
    assert failed and all("=" in f["findings"] for f in failed)
    summary_header, summary = tables["review_summary"]
    assert int(dict(zip(summary_header, summary[0]))["total"]) == len(rows)
    tables_to_xlsx(str(tmp_path / "review.xlsx"), table_sheets(out))


def test_tabulate_gives_the_cli_csv_tables(client, tmp_path):
    out = client.tabulate(check=client.check_design(uri=DESIGN), order_by=["rule", "subjects"])
    tables = {t.name: t for t in out.tables}
    assert list(tables) == ["findings", "verdicts", "verdicts_by_rule"]
    header, rows = to_rows(tables["verdicts"])
    # The header `agni check --verdicts --format csv` publishes, and its kind:ref subject spelling.
    assert header == ["verdict_id", "url", "rule", "outcome", "subjects", "statement", "context", "terms", "reason"]
    vs = [dict(zip(header, r)) for r in rows]
    passes = [v for v in vs if v["outcome"] == "pass"]
    assert passes and all(v["statement"] or v["context"] for v in passes)
    assert all(":" in v["subjects"] for v in vs)
    rules = [v["rule"] for v in vs]
    assert rules == sorted(rules) and len(set(rules)) > 1, "order_by rule did not order the verdicts"
    count_header, count_rows = to_rows(tables["verdicts_by_rule"])
    assert sum(int(dict(zip(count_header, r))["total"]) for r in count_rows) == len(rows)
    tables_to_xlsx(str(tmp_path / "check.xlsx"), table_sheets(out))


def test_tabulate_refuses_an_unknown_order_column(client):
    with pytest.raises(Exception, match="nope"):
        client.tabulate(check=client.check_design(uri=DESIGN), order_by=["nope"])


# The table rows an answer becomes are the engine's (agni issue 862). These are the public functions
# that still take an answer message, each for its stated reason, so a new one fails here.
STILL_TAKES_AN_ANSWER = {
    "rows_as_dicts": "reads a query answer as dicts; it lays no table out",
}


def test_no_other_function_lays_out_an_answer():
    import inspect

    import agni.tables as tables

    takers = set()
    for name, fn in inspect.getmembers(tables, inspect.isfunction):
        if name.startswith("_") or fn.__module__ != tables.__name__:
            continue
        hints = " ".join(str(p.annotation) for p in inspect.signature(fn).parameters.values())
        if "_pb2." in hints and "tables_pb2." not in hints:
            takers.add(name)
    assert takers, "the sweep matched nothing, so it is not reading the annotations"
    assert takers == set(STILL_TAKES_AN_ANSWER), f"functions taking an answer message: {sorted(takers)}"
