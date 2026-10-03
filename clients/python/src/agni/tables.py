"""Answers as rows, and rows as a workbook.

A query answer is ``repeated QueryRow { repeated string cells }`` plus ``column_kinds``, so its cells
are strings and stay strings here. The kinds say what a column NAMES (a net, a component), not what
type its values are, and converting ``"10k"`` or ``"R12"`` on a guess would be wrong half the time.
"""

from __future__ import annotations

import re
from typing import Dict, Iterable, List, Sequence, Tuple, Union

from agni.v1.checks import checks_pb2 as checks
from agni.v1.webapi import checks_pb2, diff_pb2, query_pb2, review_pb2

Rows = Tuple[List[str], List[List[str]]]
Table = Union[query_pb2.RunQueryResponse, checks_pb2.CheckDesignResponse, Rows]

FINDING_COLUMNS = ["rule", "severity", "subject_kind", "subject", "message", "inconclusive", "sheets"]


def rows_as_dicts(resp: query_pb2.RunQueryResponse) -> List[Dict[str, str]]:
    """Each row as ``{column: cell}``. Cells are the strings agni sent."""
    return [dict(zip(resp.columns, row.cells)) for row in resp.rows]


def to_rows(table: Table) -> Rows:
    """A header and its rows, from a query answer, a check response's findings, or a pair already.

    A check response flattens one finding per row over ``FINDING_COLUMNS``. Verdicts and skipped
    rules are not rows of that table and are left out; flatten them yourself if you want them.
    """
    if isinstance(table, query_pb2.RunQueryResponse):
        return list(table.columns), [list(r.cells) for r in table.rows]
    if isinstance(table, checks_pb2.CheckDesignResponse):
        rows = [
            [
                f.rule,
                f.severity,
                f.subject.kind,
                f.subject.ref,
                f.message,
                "yes" if f.inconclusive else "",
                " ".join(f.sheets),
            ]
            for f in table.findings
        ]
        return list(FINDING_COLUMNS), rows
    if isinstance(table, tuple) and len(table) == 2:
        header, rows = table
        return [str(h) for h in header], [[str(c) for c in r] for r in rows]
    raise TypeError(f"cannot tabulate a {type(table).__name__}")


def set_sheets(resp: query_pb2.RunQueriesResponse, allow_missing: bool = False) -> List[Tuple[str, Table]]:
    """One ``(name, answer)`` per query of a set, in the set's order, ready for ``tables_to_xlsx``.

    A query the set could not answer has no table, and a workbook that silently lacks its tab reads
    as a question that matched nothing. So an unanswered query raises ``ValueError`` naming it and its
    error, unless ``allow_missing`` asks for the answered ones alone. Names are used as sheet names,
    so keep them within Excel's 31 characters or rename the pairs before writing.
    """
    missing = [f"{r.name}: {r.error}" for r in resp.results if r.error]
    if missing and not allow_missing:
        raise ValueError("unanswered queries: " + "; ".join(missing))
    return [(r.name, r.result) for r in resp.results if not r.error]


_DIGITS = re.compile(r"(\d+)")


def natural_key(s: str) -> tuple:
    """A sort key that reads the digits in a name as a number, so ``R2`` sorts before ``R10`` and
    ``U10.3`` before ``U10.12``. Text compares case-insensitively, with the name itself breaking ties,
    so ``R01`` and ``R1`` still sort in a fixed order."""
    parts = _DIGITS.split(s)
    return tuple((int(p), p) if i % 2 else (p.casefold(), p) for i, p in enumerate(parts))


def natural_sort(table: Table, *columns: str) -> Rows:
    """The table's rows ordered by the named columns in natural order (``natural_key``), the first
    column when none is named. Every row is kept, and the header is unchanged."""
    header, rows = to_rows(table)
    idx = [header.index(c) for c in columns] or [0]
    return header, sorted(rows, key=lambda r: tuple(natural_key(r[i]) for i in idx))


# One sheet per net change kind, in the order a reviewer reads a diff. "equal" only appears when the
# request set include_equal, so its sheet is written only when it has rows: an empty "Unchanged nets"
# would read as every net having changed.
_NET_SHEETS = [
    ("new", "New nets"),
    ("deleted", "Deleted nets"),
    ("renamed", "Renamed nets"),
    ("renamed-approx", "Near renames"),
    ("hard", "Changed nets"),
    ("soft", "Attribute changes"),
    ("equal", "Unchanged nets"),
]
NET_CHANGE_COLUMNS = ["net", "old_name", "added", "removed", "old_source", "new_source"]
RENAME_EVIDENCE_COLUMNS = [
    "old_coverage", "old_coverage_significant", "new_coverage_significant", "overlap",
    "overlap_significant", "old_endpoints", "new_endpoints", "old_significant", "new_significant",
]
COMPONENT_CHANGE_COLUMNS = ["ref_des", "change", "field", "old", "new"]


def _source(prov) -> str:
    if not prov.source_file:
        return ""
    return f"{prov.source_file}:{prov.span.line}" if prov.span.line else prov.source_file


def _endpoints(eps: Iterable[str]) -> str:
    return " ".join(sorted(eps, key=natural_key))


def _number(v: float) -> str:
    return f"{v:g}"


def diff_sheets(resp: diff_pb2.DiffDesignsResponse) -> List[Tuple[str, Rows]]:
    """A diff as sheets, one per kind of net change and one for components, each in natural order.

    Every kind's sheet is written even when empty, so a workbook's tabs are the same from one run to
    the next and an empty tab says that kind did not happen. "Unchanged nets" is the exception, written
    only when the response has some (see ``_NET_SHEETS``). A near rename carries the evidence the
    engine matched it on, since it was assigned rather than recovered. A net kind this client does not
    know raises ``ValueError`` rather than being left off every sheet.
    """
    known = {k for k, _ in _NET_SHEETS}
    unknown = sorted({n.kind for n in resp.report.nets} - known)
    if unknown:
        raise ValueError(f"diff_sheets does not know the net change kind(s) {', '.join(unknown)}")
    sheets: List[Tuple[str, Rows]] = []
    for kind, name in _NET_SHEETS:
        rows = []
        for n in resp.report.nets:
            if n.kind != kind:
                continue
            row = [n.name, n.old_name, _endpoints(n.added), _endpoints(n.removed), _source(n.old_prov), _source(n.new_prov)]
            if kind == "renamed-approx":
                row += [_number(getattr(n.approx, c)) for c in RENAME_EVIDENCE_COLUMNS]
            rows.append(row)
        if kind == "equal" and not rows:
            continue
        header = NET_CHANGE_COLUMNS + (RENAME_EVIDENCE_COLUMNS if kind == "renamed-approx" else [])
        sheets.append((name, natural_sort((list(header), rows), "net")))
    r = resp.report
    comps = [[c, "added", "", "", ""] for c in r.components_added]
    comps += [[c, "removed", "", "", ""] for c in r.components_removed]
    comps += [[c.ref_des, "changed", c.field, c.old, c.new] for c in r.components_changed]
    sheets.append(("Component changes", natural_sort((list(COMPONENT_CHANGE_COLUMNS), comps), "ref_des", "change", "field")))
    return sheets


def _subject(s: checks.Subject) -> str:
    ref = f"{s.ref}.{s.pin}" if s.pin else s.ref
    return f"{s.kind} {ref}" if s.kind else ref


REVIEW_COLUMNS = ["area", "id", "title", "outcome", "note", "findings", "unmet"]


def review_sheet(review: review_pb2.Review) -> List[Tuple[str, Rows]]:
    """A review as two sheets: one row per checklist item in the checklist's own order, and the summary.

    An item's findings are written as ``rule: kind ref``, one per firing, so a failed item names what
    failed. Unmet datasheet dependencies name the part, since a reader fixes those by seeding its
    spec. Rows keep the checklist's order rather than sorting, because a reviewer reads a checklist
    in the order it was written.
    """
    rows = []
    for area in review.results.areas:
        for it in area.items:
            findings = "; ".join(f"{f.rule}: {_subject(f.subject)}" for f in it.findings)
            unmet = "; ".join(
                " ".join(x for x in (u.manufacturer, u.mpn) if x) + (" (no spec)" if u.spec_absent else "") for u in it.unmet
            )
            rows.append([area.name, it.id, it.title, it.outcome, it.note, findings, unmet])
    s = review.summary
    summary = [[k, str(getattr(s, k))] for k in ("total", "covered", "answered", "pass", "fail", "provisional")]
    return [("Review", (list(REVIEW_COLUMNS), rows)), ("Review summary", (["metric", "value"], summary))]


VERDICT_COLUMNS = ["rule", "outcome", "subject", "reason", "witness", "terms", "datasheets", "context", "id"]
_OUTCOMES = ["pass", "fail", "inconclusive", "no-limit", "not-considered"]


def _outcome(o: int) -> str:
    return checks.Outcome.Name(o).removeprefix("OUTCOME_").lower().replace("_", "-")


def _citation(d: checks.DatasheetCitation) -> str:
    return f"{d.doc or d.doc_ref} p.{d.page}" if d.page else (d.doc or d.doc_ref)


def verdict_sheets(resp: checks_pb2.CheckDesignResponse) -> List[Tuple[str, Rows]]:
    """A check run's considered set as two sheets: one row per verdict, passes included, and a count of
    each outcome per rule.

    A verdict row carries its subject, the reason, the witness that proves a pass and its terms, the
    datasheet citations and the context entities, so a passing row says why it passed. Rows are in
    rule order, then natural subject order.

    A response with findings and no verdicts was stripped of its considered set, which is what `agni
    check --format json` prints without ``--verdicts``. That raises ``ValueError``, since an empty
    sheet there would read as a run that concluded nothing.
    """
    if not resp.verdicts and resp.findings:
        raise ValueError("the response has findings and no verdicts, so its considered set was stripped; ask with --verdicts")
    rows = []
    counts: Dict[str, Dict[str, int]] = {}
    for v in resp.verdicts:
        out = _outcome(v.outcome)
        counts.setdefault(v.rule, {})[out] = counts.setdefault(v.rule, {}).get(out, 0) + 1
        rows.append([
            v.rule,
            out,
            "; ".join(_subject(s) for s in v.subjects),
            v.reason,
            v.witness.statement,
            "; ".join(f"{t.label}={t.value}" for t in v.witness.terms),
            "; ".join(_citation(d) for d in v.witness.datasheet),
            "; ".join(f"{c.role}: {_subject(c.subject)}" for c in v.context),
            v.id,
        ])
    by_rule = []
    for rule in sorted(counts, key=natural_key):
        c = counts[rule]
        by_rule.append([rule] + [str(c.get(o, 0)) for o in _OUTCOMES] + [str(sum(c.values()))])
    return [
        ("Verdicts", natural_sort((list(VERDICT_COLUMNS), rows), "rule", "subject")),
        ("Verdicts by rule", (["rule"] + _OUTCOMES + ["total"], by_rule)),
    ]


# Excel refuses these in a sheet name and caps the name at 31 characters.
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")


def tables_to_xlsx(path: str, sheets: Sequence[Tuple[str, Table]]) -> None:
    """Write one sheet per ``(name, table)``: a bold header row, frozen, with an autofilter.

    Needs the ``xlsx`` extra (``pip install agni[xlsx]``). Styling beyond that belongs to the caller,
    which can reopen the file with openpyxl. A sheet name Excel would reject raises ``ValueError``
    rather than being quietly rewritten, since a caller referring to a tab by name would then miss it.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover - exercised only without the extra
        raise ImportError("tables_to_xlsx needs openpyxl: pip install 'agni[xlsx]'") from e

    _check_names([name for name, _ in sheets])
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    for name, table in sheets:
        header, rows = to_rows(table)
        ws = wb.create_sheet(title=name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A2"
        if header:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    wb.save(path)


def _check_names(names: Iterable[str]) -> None:
    seen = set()
    for n in names:
        if not n or len(n) > 31 or _BAD_SHEET.search(n):
            raise ValueError(f"sheet name {n!r}: Excel needs 1-31 characters and none of []:*?/\\")
        if n.lower() in seen:
            raise ValueError(f"sheet name {n!r} is used twice (Excel compares names case-insensitively)")
        seen.add(n.lower())
